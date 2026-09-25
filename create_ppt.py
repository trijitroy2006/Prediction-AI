from pptx import Presentation
from pptx.util import Inches, Pt

def create_presentation():
    prs = Presentation()

    # Slide 1: Title Slide
    slide_layout = prs.slide_layouts[0] # Title slide
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Prediction AI"
    subtitle.text = "Startup & Project Risk Analyzer\nMilestone 1 Update"

    # Slide 2: Project Overview
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Project Overview"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Objective: An intelligent decision-support platform to analyze startup ideas and project proposals."
    p = tf.add_paragraph()
    p.text = "Milestone 1 Focus: Data Collection & Market Intelligence"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Target Audience: Entrepreneurs, students, startups, and project teams."
    p.level = 1

    # Slide 3: Milestone 1 Achievements
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Milestone 1 Achievements"
    tf = slide.placeholders[1].text_frame
    tf.text = "Successfully completed the core data and intelligence layer:"
    p = tf.add_paragraph()
    p.text = "Interactive UI: 3-column dashboard with smart visibility states."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Dynamic Market Analysis: Automatic calculation of TAM, SAM, SOM and growth trends based on industry."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Competitor Landscape: Real-time generation of Direct vs Indirect competitors with market share."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Robust Backend: PostgreSQL integration with seamless session fallback."
    p.level = 1

    # Slide 4: Technical Stack
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Technical Stack"
    tf = slide.placeholders[1].text_frame
    tf.text = "Technologies utilized in Milestone 1:"
    p = tf.add_paragraph()
    p.text = "Frontend: HTML5, CSS3, Jinja2 Templates (Custom Mac-style UI)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Backend Application: Python, Flask Framework"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Database: PostgreSQL (using psycopg2 for connection)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Version Control: Git & GitHub"
    p.level = 1

    # Slide 5: Next Steps
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Next Steps (Milestones 2-4)"
    tf = slide.placeholders[1].text_frame
    tf.text = "Upcoming features to be implemented:"
    p = tf.add_paragraph()
    p.text = "Milestone 2: Risk scoring engine and automated SWOT analysis."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Milestone 3: Gemini/OpenAI integration for strategic reasoning and mitigation suggestions."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Milestone 4: Comprehensive PDF assessment reports and final deployment."
    p.level = 1

    # Save the presentation
    file_name = "Milestone_1_Presentation.pptx"
    prs.save(file_name)
    print(f"Presentation saved successfully as {file_name}")

if __name__ == '__main__':
    create_presentation()
